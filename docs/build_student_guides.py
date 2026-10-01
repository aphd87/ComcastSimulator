"""Builds the two MBA student guides (TV/Streaming and Movies) as .docx.

Run:  python docs/build_student_guides.py
Output: docs/The_Slate_Student_Guide_TV_Streaming.docx
        docs/The_Slate_Student_Guide_Movies.docx

Content is written against the app as of 2026-10-01. If a section, button
label, score weight, or pass threshold changes in the app, update it here
and re-run, so the guides never drift from what students actually see.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path(__file__).parent

INK = RGBColor(0x1F, 0x23, 0x2B)
MUTED = RGBColor(0x5A, 0x60, 0x6E)
ACCENT = RGBColor(0x1A, 0x6B, 0xB5)

CALLOUT_FILL = {"tip": "E8F1FA", "warn": "FDF0E6", "key": "EEF6EC"}
CALLOUT_LABEL = {"tip": "TIP", "warn": "WATCH OUT", "key": "KEY IDEA"}


# ── helpers ──────────────────────────────────────────────────────────────────
def new_doc(footer_text: str = "The Slate — Media Portfolio Simulation · Student Guide") -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(1))
    sec.top_margin = sec.bottom_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT if name != "Heading 3" else INK
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(4)
        # python-docx headings default to a theme font; force Calibri everywhere
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), "Calibri")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(footer_text)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    return doc


def add_rich(p, text: str):
    """**bold** spans inside a plain string."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part:
            p.add_run(part).bold = (i % 2 == 1)
    return p


def para(doc, text: str, muted=False, size=None):
    p = add_rich(doc.add_paragraph(), text)
    for r in p.runs:
        if muted:
            r.font.color.rgb = MUTED
        if size:
            r.font.size = Pt(size)
    return p


def bullets(doc, items):
    for it in items:
        add_rich(doc.add_paragraph(style="List Bullet"), it)


def steps(doc, items):
    """Numbered steps. Restarts at 1 for each call by using manual numbers,
    which keeps every list independent without a numbering-definition dance."""
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        n = p.add_run(f"{i}.  ")
        n.bold = True
        n.font.color.rgb = ACCENT
        add_rich(p, it)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def callout(doc, kind: str, text: str):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    c.width = Inches(6.5)
    shade(c, CALLOUT_FILL[kind])
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    lab = p.add_run(CALLOUT_LABEL[kind] + ":  ")
    lab.bold = True
    lab.font.size = Pt(10)
    add_rich(p, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.width = Inches(widths[i])
        shade(c, "1F232B")
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Inches(widths[i])
            p = add_rich(cells[i].paragraphs[0], str(val))
            for r in p.runs:
                r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def title_block(doc, title, subtitle, label="THE SLATE  ·  STUDENT GUIDE"):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(26)
    p.paragraph_format.space_after = Pt(2)
    para(doc, subtitle, muted=True, size=12)


def page_break(doc):
    doc.add_page_break()


# ── shared sections ──────────────────────────────────────────────────────────
def sign_in_section(doc, game_button: str, movies: bool):
    doc.add_heading("Step 1 — Sign in your team", level=1)
    para(doc, "Your instructor will give you the link to the simulation. Open it in Chrome, Edge, or Safari on a laptop (not a phone).")
    steps(doc, [
        "Fill in **University**, **School / College**, and **Team Name**.",
        "Fill in **Class Title** and **Semester / Year**. Class Abbreviation is optional.",
        "Choose your role: **Driver** (makes the decisions) or **Follow Along** (view only).",
        "Click **Register Team →**.",
        f"On the next screen, click **{game_button}**.",
    ])
    callout(doc, "warn",
            "Every teammate must type the University, School, Class, and Team Name **exactly the same way**. "
            "That is how the app knows you are one team and puts you on the same leaderboard. "
            "\"Team Alpha\" and \"team alpha \" are two different teams.")
    callout(doc, "key", "Use a team name only. **Never enter student names or ID numbers.**")
    if movies:
        callout(doc, "tip",
                "The Movies simulation does not support Follow Along yet. Have **one person drive on one "
                "laptop** (project it or crowd around it), and have the others take notes and argue the numbers.")
    else:
        callout(doc, "tip",
                "Playing on several laptops? Exactly **one** person registers as Driver. Everyone else registers "
                "as Follow Along with the same details, and their screens update automatically every few seconds "
                "to show the Driver's choices.")
    callout(doc, "warn",
            "Don't close or refresh your browser tab mid-level. Your in-progress decisions live in that tab. "
            "Only submitted scores are saved to the leaderboard.")


def submit_section(doc, final_button: str, unit: str = "year", whole: str = "level"):
    doc.add_heading("Submitting your score", level=1)
    para(doc, f"After your last {unit}, click **{final_button}**. You will see your score, a breakdown, and your highlights.")
    steps(doc, [
        "Review the Score Breakdown. It shows exactly where you earned and lost points.",
        "Click **🎯 Submit Official Score**.",
        "Check the **🏆 Leaderboard** to see where you rank in your class.",
    ])
    callout(doc, "warn",
            "**Your FIRST submission is your official score.** You get 3 attempts in total, but attempts 2 and 3 "
            "are practice only. Don't click Submit until your team agrees you are done.")
    callout(doc, "tip",
            f"Before you submit, you can still go back: **← Redo** replays only the most recent {unit}, and "
            f"**↺ Restart** wipes the whole {whole}. Neither one uses up an attempt. Only Submit does.")


# ── TV / Streaming guide ─────────────────────────────────────────────────────
def build_tv():
    doc = new_doc()
    title_block(doc, "TV / Streaming",
                "Run three NBCUniversal networks — Oxygen, Bravo, and Peacock — as General Manager.")

    doc.add_heading("Quick start (read this first)", level=1)
    table(doc, ["", ""], [
        ["**Your job**", "Decide which shows to keep, cancel, and launch each year, and how much to spend on marketing, so your network makes money."],
        ["**The goal**", "Hit your network's **operating cash flow (OCF) margin** target by the end of the level. Margin = OCF ÷ revenue."],
        ["**How it's played**", "3 levels (one per network). Each level lasts **4 in-game years** (your instructor may change this). Each year = one page of decisions, then one click to simulate."],
        ["**Time needed**", "About 15–25 min for Oxygen, 20–35 min for Bravo, 25–40 min for Peacock."],
        ["**What counts**", "Only your **first submitted score** per network. Practice runs don't count."],
    ], [1.5, 5.0])
    t = doc.tables[-1]
    t.rows[0]._tr.getparent().remove(t.rows[0]._tr)  # quick-start table needs no header row

    doc.add_heading("The yearly loop", level=2)
    para(doc, "Every year works the same way. Once you've done it once, you know the whole game.")
    steps(doc, [
        "**Read the recap** at the top: how last year went and how far you are from your margin target.",
        "**Scroll down and make each decision**, section by section.",
        "**Click ▶ Simulate Year** at the bottom.",
        "**Read your Results**, then click **→ Start Year [next]**.",
    ])
    callout(doc, "key",
            "Your budget is tied to how you're doing. Beat the margin target and next year's budget grows faster. "
            "Miss it and the budget shrinks. One bad year makes the next year harder.")
    callout(doc, "tip",
            "**Check your syllabus for which network is due when.** A common setup: play **Oxygen as homework** and submit "
            "it, then play **Bravo or Peacock in class**. You can open any network directly from the **Choose Your Network** screen.")

    doc.add_heading("The three networks", level=2)
    table(doc, ["Network", "Years", "Starting budget", "Pass if margin ≥", "What makes it different"], [
        ["🔍 **Oxygen**", "2012–2015", "~$95M", "**12%**", "Cheap true-crime shows, a small loyal audience. The easiest level, so learn the controls here."],
        ["🍸 **Bravo**", "2016–2019", "~$220M", "**15%**", "Reality TV with affluent viewers who bring premium ad rates. A bigger budget, but the highest bar to clear."],
        ["🦚 **Peacock**", "2020–2023", "~$150M", "**10%**", "Streaming: you spend now and get paid later through subscribers. Adds Sports Rights bidding."],
    ], [0.95, 0.85, 1.0, 0.95, 2.75])
    para(doc, "Years shown are for the default 4-year levels. You can usually open any network at any time "
              "from the network selector, but your instructor decides the order you play in.", muted=True, size=10)

    page_break(doc)
    sign_in_section(doc, "→ Start TV / Streaming", movies=False)

    doc.add_heading("Step 2 — Find your way around", level=1)
    bullets(doc, [
        "**Choose Your Network**: the first time you open TV / Streaming, pick Oxygen, Bravo, or Peacock (whichever your instructor assigned). Later, switch any time with the network row at the top of the page. Switching starts that network fresh at Year 1.",
        "**📊 Simulation tab**: where you play. You'll spend almost all of your time here.",
        "**💹 P&L / OCF tab**: your profit-and-loss statement for the current year.",
        "**📈 10-Yr Forecast tab**: an illustrative long-run view. It is not scored.",
        "**📖 Theory tab**: short explainers on the business concepts behind the game (portfolio matrix, amortization, and so on).",
        "**Jump links** at the top of each year's page take you straight to any section.",
    ])

    page_break(doc)
    doc.add_heading("Step 3 — Make your decisions (one page per year)", level=1)
    para(doc, "Sections appear in this order, top to bottom. Each one tells you whether it is required or optional.")

    doc.add_heading("Year 1 only: Starting Position", level=2)
    para(doc, "A snapshot of what you inherited: your shows, your budget, and what your margin would be if you changed nothing. "
              "**Read this first.** It tells you how big a gap you need to close.")

    doc.add_heading("1 · 💰 Financing — set your marketing budget  (required)", level=2)
    steps(doc, [
        "Look at your two revenue sources: **Ad Revenue** (depends on ratings and marketing) and **Distribution** (fees cable companies pay you). Both shrink over time as people cut the cord.",
        "Glance at the **Linear vs. Streaming** chart. TV pays you sooner, and streaming pays off later.",
        "Set the **Marketing ($M this year)** slider.",
    ])
    bullets(doc, [
        "Each $1M of marketing adds about 1.5% to ad revenue, but the gains shrink above about $16M.",
        "If you see a warning that spend is **below $3M per show**, your shows are under-marketed. Either raise the budget or cancel shows so the money goes further.",
    ])
    callout(doc, "warn",
            "Sometimes a **Mid-Year Emergency Budget Cut** appears here in red. It's real and permanent for that year. "
            "Adjust your other decisions to fit the smaller budget.")

    doc.add_heading("2 · 🔄 Renewal — keep or cut each show  (required)", level=2)
    steps(doc, [
        "**Optional: buy Research** ($2M per show). Open **🔬 Pay for Research** and click a show. You get a 1–5 star signal for how its rating will really move this year, plus sometimes the regions where it plays well.",
        "In **Your Slate**, look at each show card: its **rating** (audience size; higher means more ad revenue), its **IP Score** (franchise value, 0–100), and its projected OCF.",
        "For every show, pick **Renew** or **Cancel**, and confirm the month it premieres.",
        "In **Primetime Scheduling**, put your shows into time slots. Or click **⚡ Auto-Fill by Rating** for a sensible starting layout.",
    ])
    callout(doc, "warn",
            "Every show is set to **Renew** by default. If you don't touch a show, you keep paying for it, even if it loses money. "
            "Look at every show's projected OCF before you move on.")
    callout(doc, "tip",
            "A show that loses money but has a high IP Score may be worth keeping for its franchise value. "
            "That is the real judgment call in this section.")
    bullets(doc, [
        "**Scheduling:** the best slot can lift a show's rating by up to 15%, and the worst slot can cut it by up to 15%. Most genres do best Tue/Wed at 8 PM. **Reality and Drama** do better on weekend nights. An unscheduled show is neutral.",
    ])

    doc.add_heading("3 · 🏈 Sports Rights  (Peacock only)", level=2)
    steps(doc, [
        "See which leagues are up for auction this year (NFL, Premier League, Olympics).",
        "Enter a bid for any league you want. The app shows how your bid compares to the market price.",
        "Click **📡 Submit Sports Bids**. You are bidding blind against six rival media and tech companies.",
    ])
    callout(doc, "key",
            "Sports bring in a burst of subscribers. Whether they **stay** afterward depends on how much you spend on original shows. "
            "Sports is the hook, and your entertainment slate is what keeps people subscribed. A winning bid locks you into a contract of 5–7 years.")

    doc.add_heading("🎬 Greenlighting — launch new shows  (optional, up to 3 per year)", level=2)
    para(doc, "There are two ways to add a show:")
    bullets(doc, [
        "**Acquire a Pitched Show**: buy a ready-made concept from the marketplace. Each one lists its audience, whether it's an international format, and any brand sponsor (sponsors lower the price). "
        "Some show **🔥 Competitor interest**: a rival network is bidding, so winning it now costs **25% more**, and if you pass, there's a **50% chance** each year the rival signs it and it's gone for good.",
        "**Build your own**: name it, fill in the **Show Concept Inputs** (genre, episodes, cost per episode, expected rating), and **write a 2–4 sentence pitch** (the concept, the hook, and who it's for). "
        "You can't greenlight your own show without a pitch.",
    ])
    callout(doc, "key",
            "**3 new-show slots per year**, shared between acquiring and building. You get 3 fresh slots each new year. "
            "There are more pitches than slots, so choose. Switching networks or restarting a level undoes the shows you added in it.")
    steps(doc, [
        "Pick or build a concept (and write its pitch if it's your own).",
        "Check the **3-Year P&L Comparison** and the **Sensitivity Analysis** (what happens if the rating or the cost per episode moves).",
        "Click **🎬 Greenlight [show] for [network]**. The production cost comes out of your budget **immediately**, and your pitch is shared on your class's **📋 Pitch Board**.",
    ])
    callout(doc, "tip",
            "Streaming often looks negative in Year 1. That is expected, because it books only one-third of a show's 3-year subscriber value in the first year while paying the full cost up front. "
            "Judge streaming shows over the full 3 years.")
    doc.add_heading("👀 The Pitch Board", level=3)
    para(doc, "Every pitch you greenlight is shared with the other teams in your class, and theirs with you. "
              "Open **👀 See what other teams in your class have pitched** above Show Concept Inputs, or the full "
              "**📋 Pitch Board** on the Leaderboard page. Use it like a development meeting: what is everyone betting on?")

    doc.add_heading("📊 Why? boxes  (optional, not scored)", level=2)
    para(doc, "Under several decisions you'll find a collapsed **📊 Why? (optional)** box: genre decay curves "
              "under Renewal, scheduling and cash-flow tools right after it, and the LTV curve and marketing-ROI chart "
              "in Greenlighting. Each one says when it's useful. Open them when you're unsure about that decision.")

    doc.add_heading("▶ Simulate the year", level=2)
    para(doc, "Check **Expected This Year** (your projected margin against the target), then click **▶ Simulate Year [N] → See Results**.")

    page_break(doc)
    doc.add_heading("Step 4 — Read your results", level=1)
    bullets(doc, [
        "**Margin vs. target**: are you on track to pass?",
        "**Rating Movers**: which shows rose or fell, and what that was worth in dollars.",
        "**🏆 Emmy Buzz**: prestige for Drama, Scripted, and Comedy shows. Good for bragging rights, but it doesn't change your money or your score.",
        "**🎲 Production Risk**: surprises such as delays or cost overruns on renewed shows.",
    ])
    para(doc, "Then click **→ Start Year [next]**, or **← Redo This Year** to try that year again.")

    submit_section(doc, "→ View Final Results & Submit Score")
    para(doc, "After you submit, click **→ Advance to [next network]** to move on, with the shows you greenlit still on your roster.")

    page_break(doc)
    doc.add_heading("How you're scored", level=1)
    para(doc, "Your score is out of 100 points and has five parts. **Passing** depends only on margin: you pass if your level OCF margin is at or above the network's target.")
    table(doc, ["Part", "Weight", "What earns full points", "In plain English"], [
        ["**OCF margin**", "35%", "40% margin", "Make money. This is the biggest piece."],
        ["**Average show ROI**", "25%", "+60% average ROI", "Your shows earn back more than they cost."],
        ["**Genre mix**", "15%", "Spending spread across genres", "Don't bet everything on one genre."],
        ["**Marketing efficiency**", "15%", "$30M revenue per $1M marketing", "Every marketing dollar pulls its weight."],
        ["**Renewal quality**", "10%", "All kept shows have positive ROI", "Cut the money-losers."],
    ], [1.6, 0.7, 1.9, 2.3])

    doc.add_heading("Common mistakes", level=1)
    bullets(doc, [
        "**Leaving every show on Renew.** Cut the shows that lose money unless you have a clear franchise reason to keep them.",
        "**Spreading marketing too thin.** Under $3M per show buys very little.",
        "**Greenlighting because you can.** New shows cost money right away, and they hurt this year's margin.",
        "**Judging streaming on Year 1.** Streaming shows pay off over three years.",
        "**Submitting too early.** Your first submission is your official score.",
        "**Mismatched team details.** Teammates who type the team name differently end up on separate teams.",
    ])

    doc.add_heading("Glossary", level=1)
    table(doc, ["Term", "Meaning"], [
        ["**OCF**", "Operating cash flow: revenue minus content costs, marketing, and overhead."],
        ["**OCF margin**", "OCF ÷ revenue. This is the number you need to hit to pass."],
        ["**Amortization**", "Spreading a show's cost over the months or years it earns money, instead of booking it all at once."],
        ["**Linear**", "Traditional TV with scheduled airings, paid by ads and cable fees."],
        ["**SVOD**", "Subscription streaming (Peacock), paid by subscribers over time."],
        ["**Rating**", "A show's 18–49 demo rating: the share of U.S. adults aged 18–49 watching an average episode (1.0 ≈ 1%). Each point is worth about $7M a year in ad revenue. Typical show: 1.0–1.5. Hit: 2.0+."],
        ["**IP Score**", "A show's franchise or brand value (0–100). A high score helps a show age better."],
        ["**Distribution revenue**", "Fees cable and satellite companies pay to carry your network."],
        ["**HHI / genre mix**", "A measure of concentration. Lower means a more diversified slate."],
    ], [1.6, 4.9])

    out = OUT_DIR / "The_Slate_Student_Guide_TV_Streaming.docx"
    doc.save(out)
    return out


# ── Movies guide ─────────────────────────────────────────────────────────────
def build_movies():
    doc = new_doc()
    title_block(doc, "Movies",
                "Run Universal Pictures' film slate: one big bet per cycle, paid up front.")

    doc.add_heading("Quick start (read this first)", level=1)
    table(doc, ["", ""], [
        ["**Your job**", "Greenlight one movie per cycle, decide how much to spend making and marketing it, and choose how to release it."],
        ["**The goal**", "Build a slate with a **positive risk-adjusted NPV** (net present value). In short, your movies should be worth more than they cost, after accounting for risk."],
        ["**How it's played**", "**5 films, one per cycle.** Each cycle covers 2 years: the first year you greenlight and produce (your t = 0 investment), the second you release. That's 10 in-game years."],
        ["**Time needed**", "About 25–45 minutes."],
        ["**What counts**", "Only your **first submitted score**. Practice runs don't count."],
    ], [1.5, 5.0])
    t = doc.tables[-1]
    t.rows[0]._tr.getparent().remove(t.rows[0]._tr)

    doc.add_heading("How Movies differs from TV", level=2)
    table(doc, ["", "TV / Streaming", "Movies"], [
        ["**Bets**", "Many shows at once", "One movie per cycle"],
        ["**Cost timing**", "Spread over years (amortized)", "**All paid up front**, before you know if it works"],
        ["**Success measure**", "OCF margin", "Risk-adjusted NPV"],
        ["**Risk**", "Steady and averaged out", "Concentrated: a single flop hurts"],
    ], [1.5, 2.5, 2.5])

    doc.add_heading("The cycle loop", level=2)
    steps(doc, [
        "**Read the recap** of last cycle (from cycle 2 on).",
        "**Scroll down and make each decision**, section by section.",
        "**Run the Theatrical Simulation** (required) and respond to anything it asks.",
        "**Click ▶ Simulate → See Results**, read them, then click **→ Start [next cycle]**.",
    ])
    callout(doc, "key",
            "Your studio starts with a **$3.5B** budget. If your movies average more than $20M NPV in a cycle, next cycle's budget grows 12%. "
            "If they lose money, it shrinks 15%. Discipline now buys you room later.")

    page_break(doc)
    sign_in_section(doc, "→ Start Movies", movies=True)

    page_break(doc)
    doc.add_heading("Step 2 — Make your decisions (one page per cycle)", level=1)
    para(doc, "Sections appear in this order, top to bottom. Jump links at the top take you to any section. "
              "**Most sections before Greenlight are optional**: they're ways to get a better movie or cheaper talent, and you can skip them.")

    doc.add_heading("Distribution Pipeline — Slate Scorecard  (read only)", level=2)
    para(doc, "Shows where each of your earlier movies is now earning money (in theaters, PVOD, streaming, and so on).")

    doc.add_heading("1 · Studio Partnerships  (optional)", level=2)
    para(doc, "Sign a long-term deal with a production company. It boosts star power on every future movie in that company's specialty genre. "
              "Click **Sign** to commit. **Watch out:** companies you pass on can be signed by a rival studio and lost for good.")

    doc.add_heading("Scouted Concepts  (optional)", level=2)
    para(doc, "Your scouts bring you a few ready-made movie ideas. Click **Option This Concept** to pre-fill your Greenlight decision. "
              "Ideas you skip may be taken by a rival studio before the next cycle.")

    doc.add_heading("Film Festival Acquisitions  (optional)", level=2)
    para(doc, "Bid on finished films at Sundance, TIFF, or Cannes. Enter a bid and click **Submit Bid**. "
              "You may be outbid, and you'll learn the result next cycle.")

    doc.add_heading("2 · Greenlight the Concept  (required)", level=2)
    para(doc, "This is the core decision. It has three parts:")
    doc.add_heading("🎬 Concept", level=3)
    bullets(doc, [
        "**Working Title**: anything you like.",
        "**Logline (required)**: 1–3 sentences on who the movie is about, what they want, and what's in the way. You can't Simulate without one. Once the film is made, the logline and how the film did appear on your class's **📋 Pitch Board**.",
        "**Genre**: Action/Tentpole, Sci-Fi/Fantasy, Animated, Horror, Comedy, Drama, or Awards/Prestige. Big genres earn more overseas. Prestige genres can win awards.",
        "**Concept Type**: New IP (the neutral baseline), Sequel (bigger opening, but the boost fades with each sequel), Family/Kids, or Indie-Horror.",
        "**Source Material**: Original Screenplay, or a Book, Video Game, or TV Show adaptation. Adaptations cost extra up front for the rights.",
        "**Optional: 🔎 Pay for Research ($4M)**: a preview of how this concept is likely to perform before you commit.",
    ])
    doc.add_heading("💰 Capital", level=3)
    bullets(doc, [
        "**Production Budget ($M)**: what it costs to make the movie. **Budget buys quality.** A gold hint shows your genre's typical budget (about $150M for a tentpole, $15M for horror). Far below it, the film looks cheap and draws weaker audiences; above it, gains taper off.",
        "**P&A / Marketing Spend ($M)**: prints and advertising. With no marketing, almost nobody shows up; past a point, extra spending barely moves the opening. Finding that sweet spot is part of the job.",
        "**Star Power (0–100)**: bigger stars mean a bigger opening, and they cost money.",
        "**Planned Opening Screens** and **IMAX release**. A gold hint shows how many screens your genre's audience can fill. Booking more mostly adds empty seats.",
        "**Financing Structure**: **Self-Finance** (all the upside, all the risk). **Pre-Sale**: distributors pay a guaranteed advance about equal to your expected international take; they recoup it first and you get half of anything beyond it. That's insurance: worth it for risky films like horror, a bad trade for a likely hit. **Tax Incentive**: shoot where a government rebate lowers your cost.",
        "**AI Production Tools**: cheaper and faster, but they cap how well critics can rate the movie. An AI-tooled prestige film can be nominated for an Oscar but can't win.",
    ])
    doc.add_heading("🌎 Distribution Strategy", level=3)
    para(doc, "Choose your **Exhibitor Posture**, which sets how hard you negotiate with theater owners. Aggressive keeps more of each ticket but gets you fewer screens. Exhibitor-Friendly gives up more of each ticket but gets you more screens. Standard sits in between.")
    callout(doc, "tip",
            "Before you move on, check **Capital at Risk** and the **Projected Range** chart. Each Bear / Base / Bull bar is "
            "what the movie's revenue is worth today **minus** Capital at Risk, so every extra $10M you commit pushes all three "
            "bars down about $10M unless the movie earns it back. The chart assumes a wide release and no reviews yet. "
            "If even the Base case is negative, rethink the plan.")

    doc.add_heading("Holding Deals  (optional)", level=2)
    para(doc, "Reserve a specific actor for **next** cycle's movie. Click **Place Hold** for a single movie or **Multi-Picture** for several. "
              "A rival may already have that actor booked, and even a confirmed hold can fall through.")

    doc.add_heading("3 · Release Strategy  (required)", level=2)
    steps(doc, [
        "**Pick a Debut Season.** Summer Tentpole and Holiday open bigger. Fall/Awards opens softer but is the path to awards. Off-Peak is neutral.",
        "**Pick a release strategy** (from Cycle 3 on; Cycles 1–2 are Wide Theatrical only). **Wide Theatrical** suits tentpoles, sci-fi, and animation. **Platform / Limited** opens in about 600 theaters and expands on word of mouth; it's the strongest choice for drama and awards titles. **Day-and-Date** premieres in theaters and on Peacock together: it gives up box office, but a Peacock premiere is valuable in itself, so it can win for horror and comedy.",
        "**Click 🎬 Run Theatrical Simulation.** This is required, and it locks in how the movie does in theaters.",
        "**Pay-1 Window Licensing**: keep the first streaming window on Peacock, license it for a flat fee, or click **🏷️ Shop This Window to Competitive Bid**. You can accept **any** bid, not just the highest: each comes with a 12- or 18-month term, and a shorter term returns the movie to Peacock sooner, so a lower bid can be worth more. Don't like the offers? **🔁 Reject All & Take It Back to Market** once; the new bids usually come in lower.",
        "**Set the PVOD price** (renting or buying at home while the movie is still new).",
        "**PVOD Market Acceptance Checks**: if viewers reject your price, choose **Hold** (keep the price and accept fewer sales) or **Cut** (lower the price to win back volume).",
    ])
    callout(doc, "warn",
            "The **▶ Simulate** button stays greyed out until you've written a Logline, run the Theatrical Simulation, and answered any PVOD rejection. "
            "If you can't click it, scroll up and look for the step you skipped.")

    page_break(doc)
    doc.add_heading("Step 3 — Read your results", level=1)
    bullets(doc, [
        "**NPV for the cycle**: did this movie create value?",
        "**Critical Reception**: what critics thought, which affects long-tail library sales and awards. It is separate from box office, so a hit can be panned and a flop can be praised.",
        "**Revenue Waterfall**: the money from each window (theaters, PVOD, streaming, TV, library).",
        "**Deal Waterfall — Who Gets Paid**: how much went to theaters, talent, and financing partners before you.",
        "**Surprises**: production trouble, word-of-mouth or piracy swings, and unexpected merchandise wins.",
    ])
    para(doc, "Then click **→ Start [next cycle]**, or **← Redo** to replay that cycle.")

    submit_section(doc, "→ View Final Slate & Submit Score", unit="cycle", whole="slate")

    page_break(doc)
    doc.add_heading("How you're scored", level=1)
    para(doc, "Your score is out of 100 points and has four parts. **Passing** depends on your decisions alone: your films' risk-adjusted NPV (before luck) must average above $0, so leaving the defaults alone always fails. "
              "Each film's graded NPV is **75% your decisions** (risk-adjusted, so the bear case counts) and **25% luck** (what the film "
              "actually earned). A breakout helps and a flop hurts, but good decisions still matter most.")
    table(doc, ["Part", "Weight", "What earns full points", "In plain English"], [
        ["**Risk-adjusted NPV**", "45%", "+$200M average per movie", "Your movies create value after accounting for risk (plus a little luck)."],
        ["**Capital efficiency**", "20%", "6× revenue per marketing dollar", "You got a lot back for what you spent."],
        ["**Strategic fit**", "20%", "Genre, season, and release choices that fit together", "Your decisions make sense as a set."],
        ["**Portfolio diversification**", "15%", "A varied slate across cycles", "Don't make the same movie five times."],
    ], [1.6, 0.7, 2.0, 2.2])

    doc.add_heading("Common mistakes", level=1)
    bullets(doc, [
        "**Leaving the defaults alone.** The default film is an under-funded tentpole; a slate of them always fails.",
        "**A cheap blockbuster.** A $20M action film looks cheap and underperforms. Fund the genre properly.",
        "**Maxing out P&A or screens.** Past the genre's sweet spot you're paying for empty seats.",
        "**Spending big on marketing for a weak concept.** Marketing buys an opening, not a good movie.",
        "**Making a summer prestige drama.** Awards voters rarely remember summer releases. Use Fall/Awards.",
        "**Five sequels in a row.** Franchise fatigue is real, and your diversification score drops.",
        "**Forgetting the Theatrical Simulation.** You can't Simulate until you've run it.",
        "**Submitting too early.** Your first submission is your official score.",
    ])

    doc.add_heading("Glossary", level=1)
    table(doc, ["Term", "Meaning"], [
        ["**NPV**", "Net present value: what a movie's future cash is worth today, minus what it cost."],
        ["**Risk-adjusted NPV**", "NPV weighted across the Bear, Base, and Bull outcomes, not just the hopeful one."],
        ["**Bear / Base / Bull**", "Pessimistic, expected, and optimistic scenarios."],
        ["**P&A**", "Prints and advertising, the marketing budget for a release."],
        ["**Windowing**", "Releasing in stages (theaters, then PVOD, then streaming, then TV) to earn money from each."],
        ["**PVOD**", "Premium video on demand: renting or buying at home while the film is still new."],
        ["**Pay-1 / Pay-2**", "The first and second streaming or pay-TV licensing windows after release."],
        ["**Day-and-Date**", "Releasing in theaters and on streaming on the same day."],
        ["**Cannibalization**", "When one window steals sales from another. For example, streaming early cuts box office."],
    ], [1.6, 4.9])

    out = OUT_DIR / "The_Slate_Student_Guide_Movies.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    for path in (build_tv(), build_movies()):
        print("wrote", path)
